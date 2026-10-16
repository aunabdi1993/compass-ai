"""Turning uploaded files into text, and choosing what fits in the prompt."""

import io
from dataclasses import dataclass

ALLOWED_EXTENSIONS = ("pdf", "docx", "txt", "md")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_DOC_CHARS = 50_000  # stored per document; the rest is cut and flagged
PROMPT_BUDGET_CHARS = 80_000  # total document text sent with each message (~20k tokens)
MIN_USEFUL_CHARS = 200  # don't squeeze in a sliver of a document
MAX_DOCUMENTS = 10  # per user


class DocumentError(ValueError):
    """The file can't be used; the message is safe to show the user."""


@dataclass
class ExtractedDocument:
    name: str
    text: str
    truncated: bool


def _extension(name: str) -> str:
    return name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("That PDF is password protected. Remove the password and try again.")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise DocumentError("I couldn't read that PDF. It may be damaged.") from exc


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for bad files
        raise DocumentError("I couldn't read that Word file. It may be damaged.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(name: str, data: bytes) -> ExtractedDocument:
    ext = _extension(name)
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentError("Please upload a PDF, Word (.docx), text or Markdown file.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentError("That file is over 10 MB. Please upload a smaller one.")

    if ext == "pdf":
        text = _pdf_text(data)
    elif ext == "docx":
        text = _docx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")

    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentError(
            "I couldn't find any text in that file. If it's a scan or an image, "
            "paste the text into your profile instead."
        )
    truncated = len(text) > MAX_DOC_CHARS
    return ExtractedDocument(name=name, text=text[:MAX_DOC_CHARS], truncated=truncated)


def select_for_prompt(documents: list[dict], budget: int = PROMPT_BUDGET_CHARS) -> tuple[list[dict], list[str]]:
    """Pick enabled documents, newest first, until the budget is spent.

    Returns (included, left_out_names). The last included document may be cut to fit;
    its dict gets "clipped": True. `documents` rows need name, content, enabled, created_at.
    """
    enabled = [d for d in documents if d.get("enabled", True)]
    enabled.sort(key=lambda d: str(d.get("created_at") or ""), reverse=True)
    included, left_out, remaining = [], [], budget
    for doc in enabled:
        content = doc["content"]
        if len(content) <= remaining:
            included.append({**doc, "clipped": False})
            remaining -= len(content)
        elif remaining >= MIN_USEFUL_CHARS:
            included.append({**doc, "content": content[:remaining], "clipped": True})
            remaining = 0
        else:
            left_out.append(doc["name"])
    return included, left_out


def budget_report(documents: list[dict], budget: int = PROMPT_BUDGET_CHARS) -> dict:
    """What the Documents page shows: how much of the context is used and what didn't fit."""
    included, left_out = select_for_prompt(documents, budget)
    return {
        "used": sum(len(d["content"]) for d in included),
        "limit": budget,
        "left_out": left_out,
        "clipped": [d["name"] for d in included if d["clipped"]],
    }
